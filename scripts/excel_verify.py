# 내보낸 엑셀을 실제 Microsoft Excel로 전체 재계산 → 앱 엔진(HyperFormula)이 기록한 값과 셀 단위 비교
# 사용: python scripts/excel_verify.py <export.xlsx>
import os, sys, shutil, tempfile
import openpyxl
import win32com.client as win32

src = os.path.abspath(sys.argv[1])
tmpdir = tempfile.mkdtemp()
recalced = os.path.join(tmpdir, 'recalc.xlsx')

app_vals = openpyxl.load_workbook(src, data_only=True)   # ExcelJS가 기록한 엔진 결과(캐시)
formulas = openpyxl.load_workbook(src, data_only=False)

xl = win32.DispatchEx('Excel.Application')
xl.Visible = False
xl.DisplayAlerts = False
try:
    wb = xl.Workbooks.Open(src, UpdateLinks=0, ReadOnly=True)
    xl.CalculateFull()
    wb.SaveCopyAs(recalced)
    wb.Close(False)
finally:
    xl.Quit()

xl_vals = openpyxl.load_workbook(recalced, data_only=True)
total = mismatch = 0
for ws in formulas.worksheets:
    a, b = app_vals[ws.title], xl_vals[ws.title]
    n = 0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                total += 1; n += 1
                va, vb = a[c.coordinate].value, b[c.coordinate].value
                same = (va == vb) or (isinstance(va, (int, float)) and isinstance(vb, (int, float)) and abs(va - vb) < 1e-6 * max(1, abs(vb))) \
                    or (va in (None, '') and vb in (None, '', 0)) \
                    or (hasattr(va, 'timestamp') and hasattr(vb, 'timestamp') and abs((va - vb).total_seconds()) < 1)
                if not same:
                    mismatch += 1
                    if mismatch <= 15:
                        print(f'MISMATCH {ws.title}!{c.coordinate}: app={va!r} excel={vb!r}  {c.value[:90]}')
    print(f'{ws.title}: {n} formulas')
print(f'TOTAL formulas={total} mismatches={mismatch}')
shutil.rmtree(tmpdir, ignore_errors=True)
sys.exit(1 if mismatch else 0)
